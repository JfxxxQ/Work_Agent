from docx import Document
from langchain.tools import tool
from pydantic import BaseModel, Field
from pathlib import Path
import time

# 读取文档入参
class ReadDocParams(BaseModel):
    file_path: str = Field(..., description="原始word文档路径")

# 写入文档入参
class WriteDocParams(BaseModel):
    output_title: str = Field(..., description="生成文档标题")
    summary_content: str = Field(..., description="已经生成好的摘要文本内容")

@tool(args_schema=ReadDocParams, description="读取本地word文档，返回全部文档文本内容")
def read_doc_tool(file_path: str) -> str:
    """读取本地word文档，返回全部文档文本内容"""
    path = Path(file_path)
    if not path.exists():
        return f"错误：文件{file_path}不存在"
    doc = Document(file_path)
    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    # 读取表格
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                row_text.append(cell.text)
            full_text.append("\t".join(row_text))
    return "\n".join(full_text)


@tool(args_schema=WriteDocParams, description="接收标题和摘要文本，生成保存新的word文档")
def write_doc_tool(output_title: str, summary_content: str) -> str:
    """接收标题和摘要文本，生成保存新的word文档，返回保存路径"""
    out_dir = Path("./word_output")
    out_dir.mkdir(exist_ok=True)
    new_doc = Document()
    new_doc.add_heading(output_title, level=1)
    new_doc.add_paragraph(summary_content)
    file_name = time.strftime("%Y%m%d%H%M%S", time.localtime()) + ".docx"
    save_path = str(out_dir / file_name)
    new_doc.save(save_path)
    return f"文档生成完成，保存路径：{save_path}"
