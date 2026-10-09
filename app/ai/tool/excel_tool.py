from pathlib import Path
import time
import openpyxl
from docx import Document
from dotenv import load_dotenv
from openpyxl import load_workbook
import re
import unicodedata
from langchain_core.tools import tool
from pydantic import BaseModel, Field
from copy import copy
import os

load_dotenv()


class ExcelParams(BaseModel):
    """工具入参模型，供LangChain Agent调用时自动校验参数"""
    excel_path: str = Field(..., description="Excel模板文件路径（第一行是表头，必须包含：区域、村民委员会(个)、总户数(户)、总人口(人)、粮食面积(亩)、粮食单产(公斤/亩)、粮食总产量(吨)）")
    word_path: str = Field(..., description="Word源文档路径，包含多条乡镇数据，每行一条记录")
    output_path: str = Field(..., description="输出的Excel文件路径，会自动创建父目录")
    field_synonyms: dict[str, list[str]] | None = Field(None, description="字段同义词映射，默认已适配模板表头，可按需扩展")
    start_row: int = Field(2, description="数据起始填入行，默认第2行（第1行是表头）")


def display_width(text: str) -> int:
    """计算文本显示宽度，中文/全角字符算2个宽度，其余算1个"""
    return sum(2 if unicodedata.east_asian_width(c) in ('F', 'W') else 1 for c in str(text or ''))


def auto_fit_columns(ws, min_w: int = 8, max_w: int = 50, padding: int = 3):
    """自动调整列宽，适配内容长度，遵循Excel可视化规范"""
    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        max_width = 0
        for c in col_cells:
            if isinstance(c, openpyxl.cell.cell.MergedCell):
                continue
            if c.value is None:
                continue
            current_width = display_width(c.value)
            if current_width > max_width:
                max_width = current_width
        final_width = max(min_w, min(max_width * 1.1 + padding, max_w))
        ws.column_dimensions[letter].width = final_width


def extract_single_row(text: str, headers: list[str], synonym_dict: dict[str, list[str]]) -> dict[str, str]:
    """
    从单条文本中提取所有字段数据
    """
    result = {}
    # 1. 特殊处理区域字段（句子开头的地名）
    if "区域" in headers:
        # 匹配句子开头的地名，直到第一个逗号
        area_match = re.search(r"^(.+?)，", text)
        if area_match:
            result["区域"] = area_match.group(1).strip()

    # 2. 处理其他常规字段
    for header in headers:
        if header == "区域":
            continue
        # 获取该字段的所有同义词
        keywords = [header]
        if header in synonym_dict:
            keywords.extend(synonym_dict[header])

        # 正则匹配：关键词后面跟空格，捕获后面的数值，直到逗号/句号
        for key in keywords:
            escaped_key = re.escape(key)
            pattern = rf"{escaped_key}\s+([^，。；\n]+)"
            match = re.search(pattern, text)
            if match:
                # 提取值，清洗单位
                raw_value = match.group(1).strip()
                # 提取纯数字（去掉个、户、人、亩、公斤每亩、吨等单位）
                num_match = re.search(r"^(\d+)", raw_value)
                if num_match:
                    result[header] = num_match.group(1)
                else:
                    result[header] = raw_value
                break
    return result


def read_docx(word_path: str) -> list[str]:
    """读取Word文档，按行拆分多条乡镇数据，返回每条数据的文本列表"""
    doc = Document(word_path)
    line_list = []
    for para in doc.paragraphs:
        para_text = para.text.strip()
        if not para_text:
            continue
        # 按句号拆分每条记录，过滤空行
        lines = [line.strip() for line in para_text.split("。") if line.strip()]
        line_list.extend(lines)
    # 处理Word内的表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                line_list.append(row_text)
    return line_list


@tool(args_schema=ExcelParams)
def fill_new_excel(excel_path: str, word_path: str, output_path: str, start_row: int = 2, field_synonyms: dict[str, list[str]] = None) -> str:
    """
    从Word文档提取多条数据，批量填入Excel模板表格
    """
    try:
        # 1. 文件校验
        excel_path = Path(excel_path)
        word_path = Path(word_path)
        output_path = Path(output_path)

        if not excel_path.exists():
            return f"错误：Excel模板不存在 {excel_path}"
        if not word_path.exists():
            return f"错误：Word源文件不存在 {word_path}"

        # 自动创建输出目录
        output_path.parent.mkdir(exist_ok=True)

        # 2. 读取Word数据，按行拆分多条记录
        word_lines = read_docx(str(word_path))
        # print(f"===== 读取到Word数据共 {len(word_lines)} 条 =====")
        # for i, line in enumerate(word_lines, 1):
        #     print(f"第{i}条：{line}")
        # print("=====================================")

        if not word_lines:
            return "错误：Word文档内容为空，未读取到有效数据"

        # 3. 加载Excel模板，读取表头
        wb = load_workbook(str(excel_path))
        ws = wb.active

        # 读取表头，清洗前后空格、全角空格
        headers = []
        col_index = []
        for col in range(1, ws.max_column + 1):
            cell_val = ws.cell(row=1, column=col).value
            if cell_val is None:
                continue
            # 清洗表头：去除前后空格、全角空格
            clean_header = str(cell_val).strip().replace("\u3000", "")
            headers.append(clean_header)
            col_index.append(col)

        if not headers:
            return "错误：Excel模板第一行没有有效表头"

        # print(f"===== 识别到Excel表头 =====")
        # for h, c in zip(headers, col_index):
        #     print(f"列号{c}：{repr(h)}")
        # print("==========================")

        # 4. 同义词字典处理，默认适配模板表头
        default_synonyms = {
            "村民委员会(个)": ["村民委员会"],
            "总户数(户)": ["总户数"],
            "总人口(人)": ["总人口"],
            "粮食面积(亩)": ["粮食种植面积"],
            "粮食单产(公斤/亩)": ["粮食单产"],
            "粮食总产量(吨)": ["粮食总产量"]
        }
        # 合并用户传入的同义词，用户传入的优先级更高
        synonym_dict = default_synonyms.copy()
        if field_synonyms:
            synonym_dict.update(field_synonyms)

        # 5. 采样模板数据行样式，用于新增内容样式复制
        template_style_row = 2  # 模板中已有数据的行号，用于样式复制
        template_cells = {}
        for col in range(1, ws.max_column + 1):
            src_cell = ws.cell(row=template_style_row, column=col)
            template_cells[col] = {
                "font": copy(src_cell.font),
                "fill": copy(src_cell.fill),
                "border": copy(src_cell.border),
                "alignment": copy(src_cell.alignment),
                "number_format": src_cell.number_format
            }

        # 6. 批量提取数据并填入Excel
        filled_count = 0
        for line_idx, line_text in enumerate(word_lines):
            # 提取单条数据
            row_data = extract_single_row(line_text, headers, synonym_dict)
            if not row_data:
                print(f"警告：第{line_idx + 1}条数据未提取到有效内容")
                continue

            # 计算当前填入的行号
            current_row = start_row + line_idx
            # 写入数据并复制样式
            for header, col in zip(headers, col_index):
                if header in row_data:
                    cell = ws.cell(row=current_row, column=col, value=row_data[header])
                    # 复制模板样式
                    if col in template_cells:
                        cell.font = copy(template_cells[col]["font"])
                        cell.fill = copy(template_cells[col]["fill"])
                        cell.border = copy(template_cells[col]["border"])
                        cell.alignment = copy(template_cells[col]["alignment"])
                        cell.number_format = template_cells[col]["number_format"]
            filled_count += 1
            #print(f"第{line_idx + 1}条数据填充完成，行号：{current_row}，数据：{row_data}")

        # 7. 自动调整列宽
        auto_fit_columns(ws)

        # 8. 保存文件
        wb.save(str(output_path))
        wb.close()

        return (
            f"填充完成！\n"
            f"共读取Word数据：{len(word_lines)} 条\n"
            f"成功填充Excel：{filled_count} 行\n"
            f"识别表头字段：{len(headers)} 个\n"
            f"输出文件：{output_path}"
        )

    except Exception as e:
        return f"填充失败，异常信息：{str(e)}"


if __name__ == "__main__":
    # 本地测试运行，适配你的文件路径
    result = fill_new_excel.invoke({
        "excel_path": "./模板.xlsx",
        "word_path": "./test2.docx",
        "output_path": "./out/res.xlsx",
        "start_row": 2
    })
    print(result)
