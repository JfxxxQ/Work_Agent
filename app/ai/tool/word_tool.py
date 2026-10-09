from docx import Document
from langchain.tools import tool
from pydantic import BaseModel, Field
from pathlib import Path
import time

class WordParams(BaseModel):
    file_path:str = Field(..., description = "原word文档路径")
    keywords:str = Field(default = "区域，人口总数，粮食面积，粮食单产，粮食总产量", description = "文档摘要类容")
    title:str = Field(..., description = "生成摘要文档的标题")


def read_docx(file_path:str) -> list[str]:
    """
    读取word文档
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"文件{file_path}不存在")
    # 打开word文档
    doc = Document(file_path)
    full_text = []

    # 读取段落
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)

    # 读取表格中的内容
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                row_text.append(cell.text)
            full_text.append("\t".join(row_text))
    return full_text

def key_word(important:list[str], keywords:str) -> list[str]:
    """
    关键词片段过滤，只筛选，不做摘要
    """
    key_list = [k.strip() for k in keywords.strip(",")]
    result = []
    for i in important:
        if any(k.strip() in i for k in key_list):
            result.append(i)
    return result

def get_create_doc_tool(model):
    @tool(args_schema = WordParams, description = "读取word文档，提取关键词内容，生成摘要word文档")
    def create_doc(file_path:str, keywords:str, title:str) -> str:
        """读取原始word，根据关键词筛选内容，调用模型生成摘要，输出新word文件"""
        try:
            #读取文档类容
            content = read_docx(file_path)
            important_list = key_word(content, keywords)
            if len(important_list) == 0:
                return "没有相关的关键词匹配，不能生成文档"
            raw_text = "\n".join(important_list)
            prompt = f"""
            请把下面文档内容精简摘要，围绕关键词：{keywords}。
            禁止复制原文大段，输出简短精炼摘要，概括全部关键信息。
            文档内容：
            {raw_text}
            """
            summary = model.invoke(prompt).content
            out_dir = Path("./word_output")
            out_dir.mkdir(exist_ok=True)
            #创建新文档
            new_doc = Document()
            new_doc.add_heading(title, level=1)
            new_doc.add_paragraph(summary)
            file_name = time.strftime("%Y%m%d%H%M%S", time.localtime()) + ".docx"
            save_path = str(out_dir / file_name)
            new_doc.save(save_path)
            return f"摘要成功"
        except Exception as e:
            return "出现异常"
    return create_doc